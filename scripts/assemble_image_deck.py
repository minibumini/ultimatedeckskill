#!/usr/bin/env python3
"""Assemble existing complete slide images into an image PPTX and PDF.

Requires python-pptx and Pillow. This does not generate or visually approve slides.
"""
import argparse
import io
import json
import math
from pathlib import Path
import tempfile
import zipfile

from PIL import Image, ImageOps, ImageColor
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor


def assemble(manifest_path, pptx_path, pdf_path):
    manifest_path = Path(manifest_path).resolve()
    pptx_path, pdf_path = Path(pptx_path).resolve(), Path(pdf_path).resolve()
    if pptx_path == pdf_path or pptx_path == manifest_path or pdf_path == manifest_path:
        raise ValueError('outputs must be distinct from each other and the manifest')
    for target in (pptx_path, pdf_path):
        if target.exists():
            raise FileExistsError(f'Will not overwrite {target}')
    data = json.loads(manifest_path.read_text(encoding='utf-8'))
    slides = data.get('slides')
    if not isinstance(slides, list) or not slides:
        raise ValueError('manifest.slides must be a nonempty list')
    canvas = data.get('canvas', {})
    width = float(canvas.get('width_inches', 13.333333))
    height = float(canvas.get('height_inches', 7.5))
    if not all(math.isfinite(x) and x > 0 for x in (width, height)):
        raise ValueError('canvas dimensions must be finite positive inches')
    background_color = ImageColor.getrgb(data.get('background', '#FFFFFF'))[:3]
    sources, page_images = [], []
    try:
        for index, item in enumerate(slides, 1):
            if not isinstance(item, dict) or not isinstance(item.get('image'), str):
                raise ValueError(f'slide {index}: image path required')
            path = (manifest_path.parent / item['image']).resolve()
            if path in (pptx_path, pdf_path):
                raise ValueError('output cannot overwrite input image')
            with Image.open(path) as original:
                oriented = ImageOps.exif_transpose(original)
                if not math.isclose(oriented.width / oriented.height, width / height, rel_tol=0.005):
                    raise ValueError(f'slide {index}: image aspect ratio differs from canvas: {path}')
                image = oriented.convert('RGBA')
                background = Image.new('RGBA', image.size, (*background_color, 255))
                background.alpha_composite(image)
                page_images.append(background.convert('RGB'))
            sources.append(str(path))
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(width), Inches(height)
        deck.core_properties.title = str(data.get('title', ''))
        # Build both files before publishing either, so validation errors do not
        # leave a partially assembled user-facing artifact.
        with tempfile.TemporaryDirectory(prefix='image-deck-', dir=manifest_path.parent) as temporary:
            temporary = Path(temporary)
            pdf_pages = []
            pdf_canvas = (round(width * 150), round(height * 150))
            for index, (item, page) in enumerate(zip(slides, page_images), 1):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(*background_color)
                scale = min(deck.slide_width / page.width, deck.slide_height / page.height)
                picture_width, picture_height = round(page.width * scale), round(page.height * scale)
                stream = io.BytesIO()
                page.save(stream, format='PNG')
                stream.seek(0)
                picture = slide.shapes.add_picture(stream, (deck.slide_width - picture_width) // 2,
                    (deck.slide_height - picture_height) // 2, width=picture_width, height=picture_height)
                picture.name = f"page-{index:03}: {item.get('title', '')}"
                notes = item.get('notes', '')
                if not isinstance(notes, str):
                    raise ValueError(f'slide {index}: notes must be a string')
                if notes:
                    slide.notes_slide.notes_text_frame.text = notes
                pdf_page = Image.new('RGB', pdf_canvas, background_color)
                fitted = ImageOps.contain(page, pdf_canvas, Image.Resampling.LANCZOS)
                pdf_page.paste(fitted, ((pdf_canvas[0] - fitted.width) // 2, (pdf_canvas[1] - fitted.height) // 2))
                pdf_pages.append(pdf_page)
            staged_pptx, staged_pdf = temporary / 'deck.pptx', temporary / 'deck.pdf'
            deck.save(staged_pptx)
            pdf_pages[0].save(staged_pdf, format='PDF', save_all=True,
                              append_images=pdf_pages[1:], resolution=150, quality=95)
            with zipfile.ZipFile(staged_pptx) as package:
                if damaged := package.testzip():
                    raise ValueError(f'invalid PPTX package: {damaged}')
            if len(Presentation(staged_pptx).slides) != len(slides):
                raise ValueError('assembled slide count mismatch')
            for target in (pptx_path, pdf_path):
                target.parent.mkdir(parents=True, exist_ok=True)
            # Exclusive creation avoids overwriting even if another process
            # creates the requested file after the initial existence check.
            for staged, target in ((staged_pptx, pptx_path), (staged_pdf, pdf_path)):
                with target.open('xb') as result:
                    result.write(staged.read_bytes())
    finally:
        for page in page_images:
            page.close()
    return {'slides': len(slides), 'pptx': str(pptx_path), 'pdf': str(pdf_path),
            'sources': sources, 'format': 'image slides; text is not individually editable',
            'generation_verified': False, 'visual_qa': 'NOT_RUN'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('manifest', type=Path)
    parser.add_argument('--pptx', type=Path, required=True)
    parser.add_argument('--pdf', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(assemble(args.manifest, args.pptx, args.pdf), ensure_ascii=False))


if __name__ == '__main__':
    main()
